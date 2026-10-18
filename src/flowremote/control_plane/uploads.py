"""Persist an authenticated user's workbook without inventing analysis results."""

from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook

from flowremote.control_plane.db import Database, new_id, utc_now
from flowremote.control_plane.order_validation import analyse_rows, save_analysis

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 256
MAX_HEADER_ROW = 1000


class UploadError(ValueError):
    pass


def read_workbook(data: bytes) -> tuple[str, int, tuple, list]:
    """Read the first sheet, preserving physical row numbers and raw cell values."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024
            ):
                raise UploadError(
                    "The expanded workbook is too large. Please upload a smaller file."
                )
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.worksheets:
                raise UploadError("The workbook must contain a worksheet.")
            sheet = workbook.worksheets[0]
            # Do not trust cached dimensions: some exporters write incorrect bounds.
            sheet.reset_dimensions()
            count = 0
            header_row = None
            header = ()
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if len(row) > MAX_COLUMNS:
                    raise UploadError("Use at most 256 columns in the first worksheet.")
                if header_row is None and index > MAX_HEADER_ROW:
                    raise UploadError(
                        "Place column headers within the first 1,000 rows of the first worksheet."
                    )
                if header_row is not None and index > header_row + MAX_ROWS:
                    raise UploadError("Use at most 10,000 worksheet rows below the headers.")
                populated = any(value is not None and str(value).strip() for value in row)
                if not populated:
                    continue
                if header_row is None:
                    header_row = index
                    header = row
                else:
                    count += 1
                    rows.append((index, row))
            if header_row is None:
                raise UploadError(
                    f'The first worksheet "{sheet.title}" is empty. '
                    "Add column headers and at least one data row, or move your data worksheet first."
                )
            if count == 0:
                raise UploadError(
                    f'The first worksheet "{sheet.title}" has no data below its headers. '
                    "The first non-empty row is treated as headers; add at least one data row below it."
                )
            return sheet.title, header_row, header, rows
        finally:
            workbook.close()
    except UploadError:
        raise
    except Exception as exc:
        raise UploadError(
            "This file could not be read as an Excel workbook. Save it as .xlsx and try again."
        ) from exc


def inspect_workbook(data: bytes) -> tuple[str, int]:
    worksheet, _, _, rows = read_workbook(data)
    return worksheet, len(rows)


def create_uploaded_task(
    db: Database, upload_dir: Path, owner_id: str, name: str, filename: str, data: bytes
) -> str:
    worksheet, header_row, header, rows = read_workbook(data)
    results, summary = analyse_rows(header, rows)
    count = len(results)
    task_id, now = new_id("TASK"), utc_now()
    storage_name = uuid4().hex + ".xlsx"
    upload_dir.mkdir(parents=True, exist_ok=True)
    destination = upload_dir / storage_name
    try:
        with destination.open("xb") as target:
            target.write(data)
        with db.transaction():
            db.execute(
                "INSERT INTO tasks (id, workflow, batch_size, status, source_file, source_location,"
                " primary_target, secondary_target, owner_id, created_at, updated_at)"
                " VALUES (?, ?, ?, 'ready', ?, 'Uploaded file', ?, ?, ?, ?, ?)",
                (
                    task_id,
                    name,
                    count,
                    filename,
                    "OA - Order Administration",
                    "BI - Business Intelligence",
                    owner_id,
                    now,
                    now,
                ),
            )
            db.execute(
                "INSERT INTO task_uploads (task_id, storage_name, byte_size, worksheet, created_at)"
                " VALUES (?, ?, ?, ?, ?)",
                (task_id, storage_name, len(data), worksheet, now),
            )
            save_analysis(db, task_id, header_row, results, summary)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    return task_id


def validate_saved_upload(db: Database, upload_dir: Path, task_id: str):
    rows = db.query("SELECT storage_name FROM task_uploads WHERE task_id = ?", (task_id,))
    if not rows:
        raise UploadError("This task has no uploaded workbook.")
    path = upload_dir / rows[0]["storage_name"]
    if not path.is_file():
        raise UploadError("The original workbook is missing. Please upload it as a new task.")
    _, header_row, header, records = read_workbook(path.read_bytes())
    results, summary = analyse_rows(header, records)
    with db.transaction():
        # Repeated/concurrent submissions must never erase review decisions.
        if db.scalar("SELECT status FROM tasks WHERE id = ?", (task_id,)) != "ready":
            return
        save_analysis(db, task_id, header_row, results, summary)
